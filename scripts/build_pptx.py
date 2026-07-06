"""원고(md, manuscript-schema 문법) → PPTX. 발표자 노트에 Narration 삽입.

사용: python scripts/build_pptx.py <manuscript.md> <out.pptx> [--assets-root <course-dir>]
"""
import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

# 자산 임베드 안전 여백 정책: EMBED_SAFE_MARGIN_RATIO = 0.05 (스펙 §3.0-A canonical)
# PPTX target 구현 상수 — 슬라이드 가장자리에서 확보할 안전 여백.
PPTX_EMBED_MARGIN = Inches(0.5)


def _fit_in_box(img_px_w, img_px_h, box_w_emu, box_h_emu):
    """이미지 픽셀 종횡비를 유지한 채 박스(EMU) 안에 들어가는 (w, h)를 반환.
    width·height 둘 다 박스를 넘지 않는다(오버플로 불가). codex 검증 조건 1."""
    aspect = img_px_w / img_px_h if img_px_h else 1.0
    w = box_w_emu
    h = int(round(w / aspect))
    if h > box_h_emu:
        h = box_h_emu
        w = int(round(h * aspect))
    return w, h


def _image_px_size(path):
    """PIL로 픽셀 크기 읽기(EXIF 회전 보정)."""
    from PIL import Image, ImageOps
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)
        return im.size  # (w, h)

SLIDE_RE = re.compile(r"^## Slide (\d+)\.\s*(.+)$")
FIELD_RE = re.compile(r"^\*\*(Screen|Easy analogy|Practical case|Visual asset|Source|Narration|Practice|Assessment)\*\*")
IMG_PATH_RE = re.compile(r"(assets[/\\][^\s)`\"']+\.(?:png|jpg|jpeg|webp))", re.IGNORECASE)

# Screen 필드의 최상위 라벨 줄: `- 제목: ...`, `- 짧은 문구:` 등.
SCREEN_LABEL_RE = re.compile(r"^\s*[-*]\s*([^:：]+?)\s*[:：]\s*(.*)$")
# 내용 리스트 라벨 우선순위 — 이 라벨의 들여쓴 하위 항목이 슬라이드 본문이 된다.
BODY_LIST_LABELS = ["짧은 문구", "학습목표", "학습내용"]


def _screen_title_body(screen_raw, header_title):
    """Screen 필드 원문 줄에서 표시 제목과 본문 리스트를 뽑는다.

    - 제목 = `- 제목:` 값 (없으면 슬라이드 헤더 단어로 폴백).
    - 본문 = `짧은 문구`/`학습목표`/`학습내용` 중 하위 항목이 있는 첫 라벨의 항목들.
             내용 리스트가 없으면 `부제` 값으로 폴백(표지용). 그 외 메타/연출 라벨
             (`관련 학습목표`, `핵심 정의`, `화면` 등)은 본문에서 제외한다.
    """
    title = None
    subtitle = None
    groups = {}          # label -> list of sub-items
    cur_label = None
    for raw in screen_raw:
        stripped = raw.strip()
        if not stripped:
            continue
        indent = len(raw) - len(raw.lstrip())
        m = SCREEN_LABEL_RE.match(stripped)
        # 최상위 Screen 불릿(들여쓰기 ≤1)이면서 `label: value` 꼴이면 라벨 줄.
        if indent <= 1 and m:
            label = m.group(1).strip()
            value = m.group(2).strip()
            cur_label = label
            groups.setdefault(label, [])
            if label == "제목":
                title = value
            elif label == "부제":
                subtitle = value
            continue
        # 들여쓴 하위 항목 → 현재 라벨에 귀속.
        item = re.sub(r"^[-*]\s*", "", stripped)
        item = re.sub(r"^\d+[.)]\s*", "", item)
        if cur_label and item:
            groups.setdefault(cur_label, []).append(item)

    body = []
    for lbl in BODY_LIST_LABELS:
        if groups.get(lbl):
            body = groups[lbl]
            break
    if not body and subtitle:
        body = [subtitle]
    return (title or header_title), body

def _finalize(cur):
    """screen_raw에서 표시 제목·본문을 뽑아 cur에 채운다."""
    title, body = _screen_title_body(cur["screen_raw"], cur["header_title"])
    cur["title"] = title
    cur["body_lines"] = body
    return cur


def parse_manuscript(md_text):
    """슬라이드 블록을 dict 리스트로. keys: title, body_lines, screen_lines, image_paths, code, narration"""
    slides, cur, field = [], None, None
    in_code, code_lines, fence_lang = False, [], ""
    for line in md_text.splitlines():
        if in_code:
            if line.strip().startswith("```"):
                # d2 fence = diagram source, not display code -> discard.
                if field == "Visual asset" and fence_lang != "d2" and code_lines:
                    block = "\n".join(code_lines)
                    cur["code"] = (cur["code"] + "\n\n" + block) if cur["code"] else block
                in_code, code_lines, fence_lang = False, [], ""
                continue
            code_lines.append(line)
            continue
        m = SLIDE_RE.match(line)
        if m:
            if cur:
                slides.append(_finalize(cur))
            cur = {"header_title": m.group(2).strip(), "title": m.group(2).strip(),
                   "screen_lines": [], "screen_raw": [], "body_lines": [],
                   "image_paths": [], "code": "", "narration": ""}
            field, in_code, code_lines, fence_lang = None, False, [], ""
            continue
        if cur is None:
            continue
        f = FIELD_RE.match(line)
        if f:
            field = f.group(1)
            continue
        if line.strip().startswith("```"):
            in_code = True
            fence_lang = line.strip()[3:].strip().lower()
            continue
        text = line.strip().lstrip("-").strip()
        if not text:
            continue
        if field == "Screen":
            cur["screen_lines"].append(text)
            cur["screen_raw"].append(line)
        elif field == "Narration":
            cur["narration"] = (cur["narration"] + " " + text).strip()
        elif field == "Visual asset":
            for im in IMG_PATH_RE.findall(line):
                cur["image_paths"].append(im)
    if cur:
        slides.append(_finalize(cur))
    return slides

def build_pptx(md_text, out_path, assets_root="."):
    slides = parse_manuscript(md_text)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[5]  # Title Only
    for s in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = s["title"]
        top = Inches(1.8)
        body_lines = s["body_lines"]
        if body_lines:
            box = slide.shapes.add_textbox(Inches(0.8), top, Inches(7.0), Inches(4.8))
            tf = box.text_frame
            tf.word_wrap = True
            for i, l in enumerate(body_lines[:8]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = l
                p.font.size = Pt(22)
        for im in s["image_paths"][:1]:
            p = Path(assets_root) / im
            if p.exists():
                # 이미지 배치 박스: 본문이 있으면 우측 컬럼, 없으면 하단 중앙 넓은 영역.
                # 어느 경우든 슬라이드 가장자리에서 PPTX_EMBED_MARGIN 여백 확보.
                m = PPTX_EMBED_MARGIN
                box_top = top  # 제목 아래
                box_bottom = prs.slide_height - m
                box_h = box_bottom - box_top
                if body_lines:
                    box_left = Inches(8.0)
                    box_right = prs.slide_width - m
                else:
                    box_left = m
                    box_right = prs.slide_width - m
                box_w = box_right - box_left
                try:
                    px_w, px_h = _image_px_size(p)
                except Exception:
                    px_w, px_h = (4, 3)
                w, h = _fit_in_box(px_w, px_h, box_w, box_h)
                # 박스 안에서 중앙 정렬
                left = box_left + (box_w - w) // 2
                pic_top = box_top + (box_h - h) // 2
                slide.shapes.add_picture(str(p), left, pic_top, width=w, height=h)
        if s["code"]:
            cb = slide.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.8))
            cf = cb.text_frame
            cf.word_wrap = True
            cf.text = s["code"][:600]
            cf.paragraphs[0].font.name = "Consolas"
            cf.paragraphs[0].font.size = Pt(12)
        if s["narration"]:
            slide.notes_slide.notes_text_frame.text = s["narration"]
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return len(slides)

def build_pptx_from_images(md_text, image_dir, out_path, allow_count_mismatch=False):
    """이미지 모드: ppt_preview에서 렌더한 슬라이드 PNG를 각 장의 전체 배경으로 넣고,
    원고 Narration을 발표자 노트에 삽입한다. 미리보기와 픽셀 동일한 PPTX를 만든다.

    - 이미지는 `<image_dir>/slideNN.png`(1부터, 2자리 zero-pad) 순서로 소비한다.
    - 이미지는 16:9로 렌더되어 있으므로 슬라이드(13.333x7.5") 전체에 채운다.
    - 이미지 수와 원고 슬라이드(=Narration) 수가 다르면 기본적으로 실패한다(계약 강제).
      preview/원고 drift가 산출물에 박제되는 것을 막는다. 의도적 불일치는
      allow_count_mismatch=True(CLI `--allow-count-mismatch`)로만 완화한다.
    """
    slides = parse_manuscript(md_text)
    narrations = [s["narration"] for s in slides]
    images = sorted(Path(image_dir).glob("slide*.png"))
    if not images:
        raise SystemExit(f"슬라이드 이미지 없음: {image_dir}/slide*.png")
    if len(images) != len(narrations):
        msg = (f"슬라이드 수 불일치: 렌더 이미지 {len(images)}장 vs 원고 슬라이드 "
               f"{len(narrations)}개. preview를 최신 원고로 재생성/재렌더하세요 "
               f"(의도적이면 --allow-count-mismatch).")
        if not allow_count_mismatch:
            raise SystemExit(msg)
        print("경고: " + msg, file=sys.stderr)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]  # Blank
    for i, img in enumerate(images):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
        if i < len(narrations) and narrations[i]:
            slide.notes_slide.notes_text_frame.text = narrations[i]
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return len(images)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("manuscript")
    ap.add_argument("out")
    ap.add_argument("--assets-root", default=".")
    ap.add_argument("--from-images", default=None,
                    help="이미지 모드: 이 디렉터리의 slideNN.png를 전체 배경으로 사용(ppt_preview 렌더 결과)")
    ap.add_argument("--allow-count-mismatch", action="store_true",
                    help="이미지 모드에서 렌더 이미지 수 != 원고 슬라이드 수여도 진행(기본은 실패)")
    args = ap.parse_args()
    md = Path(args.manuscript).read_text(encoding="utf-8")
    if args.from_images:
        n = build_pptx_from_images(md, args.from_images, args.out,
                                   allow_count_mismatch=args.allow_count_mismatch)
        print(f"OK (image mode): {n} slides -> {args.out}")
    else:
        n = build_pptx(md, args.out, args.assets_root)
        print(f"OK: {n} slides -> {args.out}")

if __name__ == "__main__":
    sys.exit(main())
