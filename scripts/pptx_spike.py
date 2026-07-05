"""python-pptx speaker notes spike: 슬라이드 1장 + 노트 삽입 → 재독해 검증."""
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = "scripts/spike_out.pptx"
NOTE = "안녕하세요. 이번 시간에는 스프링 부트 실행 환경을 살펴보겠습니다."

def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)   # 16:9
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "서버 프로그램과 실행 환경"
    body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(4))
    body.text_frame.text = "요청을 받는다 → 처리한다 → 응답을 돌려준다"
    body.text_frame.paragraphs[0].font.size = Pt(28)
    slide.notes_slide.notes_text_frame.text = NOTE   # 발표자 노트
    prs.save(OUT)

def verify():
    prs = Presentation(OUT)
    slide = prs.slides[0]
    assert slide.has_notes_slide, "notes_slide 없음"
    got = slide.notes_slide.notes_text_frame.text
    assert got == NOTE, f"노트 불일치: {got!r}"
    print("SPIKE PASS: speaker notes OK")

if __name__ == "__main__":
    build()
    verify()
