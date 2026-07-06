"""build_pptx 파서/빌더 테스트 (pytest)."""
import textwrap
from pptx import Presentation
from build_pptx import parse_manuscript, build_pptx

SAMPLE = textwrap.dedent("""\
    # 1차시 원고: 테스트

    ## 차시 정보
    - 과정명: 테스트 과정

    ## Slide 1. 표지
    **Screen**
    - 제목: 테스트 강의
    - 부제: 부제목입니다
    **Narration**
    - 안녕하세요. 첫 번째 나레이션입니다.

    ## Slide 2. 본문
    **Screen**
    - 제목: 본문 슬라이드
    - 짧은 문구:
      - 첫 번째 포인트
      - 두 번째 포인트
    **Visual asset**
    - GPT image prompt: `a test image`
    **Narration**
    - 두 번째 나레이션입니다.
    """)

def test_parse_slide_count_and_titles():
    slides = parse_manuscript(SAMPLE)
    assert len(slides) == 2
    assert slides[0]["title"] == "표지"
    assert slides[1]["title"] == "본문"

def test_parse_narration_and_screen():
    slides = parse_manuscript(SAMPLE)
    assert slides[0]["narration"].startswith("안녕하세요.")
    assert "첫 번째 포인트" in slides[1]["screen_lines"]

def test_build_writes_notes(tmp_path):
    out = tmp_path / "out.pptx"
    build_pptx(SAMPLE, str(out), assets_root=str(tmp_path))
    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    assert prs.slides[0].notes_slide.notes_text_frame.text.startswith("안녕하세요.")


MULTI_CODE_SAMPLE = textwrap.dedent("""\
    # 1차시 원고: 테스트

    ## Slide 1. 요청과 응답

    **Screen**
    - 제목: 요청과 응답 슬라이드

    **Visual asset**
    - Code block for slide:

    ```http
    GET /hello HTTP/1.1
    Host: localhost:8080
    ```

    ```http
    HTTP/1.1 200 OK
    Content-Type: text/plain

    Hello Spring Boot
    ```

    **Narration**
    - 나레이션입니다.
    """)

D2_ONLY_SAMPLE = textwrap.dedent("""\
    # 1차시 원고: 테스트

    ## Slide 1. HTTP 흐름

    **Screen**
    - 제목: HTTP 요청과 응답

    **Visual asset**
    - D2 diagram: `assets/diagrams/ch01_http-request-response.d2`

    ```d2
    browser: "브라우저"
    server: "서버"
    browser -> server: "요청 전달"
    ```

    **Narration**
    - 나레이션입니다.
    """)

FENCE_FALSE_POSITIVE_SAMPLE = textwrap.dedent("""\
    # 1차시 원고: 테스트

    ## Slide 1. 첫 슬라이드
    **Screen**
    - 제목: 첫 슬라이드

    **Visual asset**
    - Code block for slide:

    ```text
    ## Slide 99. 가짜 슬라이드처럼 보이는 코드 줄
    **Narration**
    이건 코드 안의 텍스트일 뿐입니다.
    ```

    **Narration**
    - 진짜 나레이션입니다.

    ## Slide 2. 두 번째 슬라이드
    **Screen**
    - 제목: 두 번째 슬라이드
    **Narration**
    - 두 번째 나레이션입니다.
    """)


def test_multiple_code_fences_are_preserved():
    slides = parse_manuscript(MULTI_CODE_SAMPLE)
    assert len(slides) == 1
    code = slides[0]["code"]
    assert "GET /hello HTTP/1.1" in code
    assert "HTTP/1.1 200 OK" in code
    assert "Hello Spring Boot" in code


def test_d2_fence_is_not_treated_as_display_code():
    slides = parse_manuscript(D2_ONLY_SAMPLE)
    assert len(slides) == 1
    assert slides[0]["code"] == ""


def test_slide_and_field_markers_inside_code_fence_are_ignored():
    slides = parse_manuscript(FENCE_FALSE_POSITIVE_SAMPLE)
    assert len(slides) == 2
    assert slides[0]["title"] == "첫 슬라이드"
    assert slides[1]["title"] == "두 번째 슬라이드"
    assert "## Slide 99." in slides[0]["code"]
    assert slides[0]["narration"] == "진짜 나레이션입니다."
    assert slides[1]["narration"] == "두 번째 나레이션입니다."


# --- 안전 여백(fit-in-box) 회귀 테스트 (2026-07-06) ---
def _make_img(tmp_path, name, w, h):
    from PIL import Image
    p = tmp_path / name
    Image.new("RGB", (w, h), "white").save(p)
    return p


def _manuscript_with_image(rel_path):
    import textwrap
    return textwrap.dedent(f"""\
        # 테스트

        ## Slide 1. 이미지
        **Screen**
        - 제목: 이미지 슬라이드
        **Visual asset**
        - GPT image prompt: `x`
        - → 생성됨: {rel_path}
        **Narration**
        - 노트.
        """)


def _picture_bounds(pptx_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(pptx_path))
    sw, sh = prs.slide_width, prs.slide_height
    pics = []
    for s in prs.slides:
        for sh_ in s.shapes:
            if sh_.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pics.append((sh_.left, sh_.top, sh_.width, sh_.height))
    return sw, sh, pics


def test_ultra_tall_image_stays_in_slide(tmp_path):
    """초세로 이미지(100x2000)가 슬라이드 경계를 넘지 않는다."""
    (tmp_path / "assets" / "images" / "ch01").mkdir(parents=True)
    _make_img(tmp_path, "assets/images/ch01/slide01.png", 100, 2000)
    out = tmp_path / "out.pptx"
    build_pptx(_manuscript_with_image("assets/images/ch01/slide01.png"), str(out), assets_root=str(tmp_path))
    sw, sh, pics = _picture_bounds(out)
    assert len(pics) == 1
    left, top, w, h = pics[0]
    assert top + h <= sh, f"세로 오버플로: bottom={top+h} > slide={sh}"
    assert left + w <= sw, f"가로 오버플로: right={left+w} > slide={sw}"


def test_ultra_wide_image_stays_in_slide(tmp_path):
    """초광폭 이미지(2000x100)가 슬라이드 경계를 넘지 않는다."""
    (tmp_path / "assets" / "images" / "ch01").mkdir(parents=True)
    _make_img(tmp_path, "assets/images/ch01/slide01.png", 2000, 100)
    out = tmp_path / "out.pptx"
    build_pptx(_manuscript_with_image("assets/images/ch01/slide01.png"), str(out), assets_root=str(tmp_path))
    sw, sh, pics = _picture_bounds(out)
    left, top, w, h = pics[0]
    assert top + h <= sh and left + w <= sw


def test_image_leaves_margin_from_slide_edges(tmp_path):
    """이미지가 슬라이드 가장자리에 닿지 않고 여백을 남긴다."""
    (tmp_path / "assets" / "images" / "ch01").mkdir(parents=True)
    _make_img(tmp_path, "assets/images/ch01/slide01.png", 800, 600)
    out = tmp_path / "out.pptx"
    build_pptx(_manuscript_with_image("assets/images/ch01/slide01.png"), str(out), assets_root=str(tmp_path))
    sw, sh, pics = _picture_bounds(out)
    left, top, w, h = pics[0]
    # 우측·하단 가장자리에서 최소 여백(0.2in = 182880 EMU)
    assert sw - (left + w) >= 182880, "우측 여백 부족"
    assert sh - (top + h) >= 182880, "하단 여백 부족"
